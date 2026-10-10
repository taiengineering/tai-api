"""
Structure and contract validation for generated .xlsx files.
WO-REF01-XLS03-BULK-XLSX-BUILD-001 Phase F.
"""
from __future__ import annotations

import os
import zipfile
from typing import Optional

import openpyxl


_FORBIDDEN_FORMULA_IDS = {"CHW-03", "REF-C029", "REF-C004"}


def validate_xlsx(path: str, research_id: str, design_type: str) -> list[str]:
    """
    Returns list of failure strings. Empty list = PASS.
    """
    failures = []

    if not os.path.exists(path):
        return [f"FILE_NOT_FOUND: {path}"]

    # 1. Valid zip (xlsx is a zip)
    if not zipfile.is_zipfile(path):
        return ["INVALID_ZIP: file is not a valid xlsx/zip"]

    try:
        wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
    except Exception as e:
        return [f"LOAD_ERROR: {e}"]

    if not wb.sheetnames:
        return ["NO_SHEETS: workbook has no sheets"]

    ws = wb.active

    # 2. Has at least 3 rows (title + basic_info + at least one data row)
    max_row = ws.max_row or 0
    if max_row < 3:
        failures.append(f"TOO_FEW_ROWS: max_row={max_row}, expected >=3")

    # 3. Has at least 1 column
    max_col = ws.max_column or 0
    if max_col < 1:
        failures.append(f"NO_COLUMNS: max_column={max_col}")

    # 4. Title cell (A1) is non-empty
    title_cell = ws.cell(row=1, column=1).value
    if not title_cell:
        failures.append("EMPTY_TITLE: cell A1 is empty")

    # 5. Auto-filter present (required for forms with repeat_table)
    # MNT-03 has no repeat_table — skip filter check for it
    if research_id not in {"MNT-03"}:
        if not ws.auto_filter.ref:
            failures.append("NO_AUTO_FILTER: auto_filter.ref not set")

    # 6. Print area set
    if not ws.print_area:
        failures.append("NO_PRINT_AREA: print_area not set")

    # 7. FORMULA_DIRECTION_UNVERIFIED: no formula cells in data area
    if research_id in _FORBIDDEN_FORMULA_IDS:
        formula_cells = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.append(cell.coordinate)
        if formula_cells:
            failures.append(
                f"FORBIDDEN_FORMULA: formulas found in {research_id}: {formula_cells[:5]}"
            )

    # 8. Freeze panes set (for forms with repeat_table)
    if research_id not in {"MNT-03"}:
        if not ws.freeze_panes:
            failures.append("NO_FREEZE_PANES: freeze_panes not set")

    # 9. Page orientation recorded
    orientation = ws.page_setup.orientation
    if orientation not in ("portrait", "landscape", None):
        failures.append(f"BAD_ORIENTATION: {orientation}")

    wb.close()
    return failures
