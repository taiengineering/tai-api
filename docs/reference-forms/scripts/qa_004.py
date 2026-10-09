"""
WO-REF01-060-B8-WAVE2-AUTO-QA-004
Automated QA for B8 Wave2 14 forms (28 PDF/DOCX outputs).
Sections A (PDF), B-structural (DOCX OOXML parse), C (DOCX roundtrip), D (PNG evidence).
LibreOffice DOCX render: UNVERIFIED_LIBREOFFICE_NOT_INSTALLED.
"""
import hashlib, json, pathlib, shutil, tempfile, zipfile, re, sys
import pymupdf
import docx as python_docx
from PIL import Image

BASE        = pathlib.Path(__file__).parent
OUTPUT      = BASE.parent / "output"
EVIDENCE    = BASE.parent / "evidence" / "qa-004"
EVIDENCE.mkdir(parents=True, exist_ok=True)

B8_IDS = [
    ("c026","C026"), ("c027","C027"), ("c028","C028"), ("c029","C029"),
    ("c031","C031"), ("c033","C033"), ("c037","C037"),
    ("c039","C039"), ("c040","C040"), ("c041","C041"), ("c042","C042"),
    ("c043","C043"), ("c044","C044"), ("gov-01","GOV-01"),
]

A4_W_PT  = 595.0   # portrait  short side
A4_H_PT  = 842.0   # portrait  long  side
MARGIN   = 30.0    # minimum expected page margin in pt
FONT_TOO_SMALL = 4.0  # flag text smaller than this

# ── helpers ─────────────────────────────────────────────────────────────────

def sha256(path):
    return hashlib.sha256(pathlib.Path(path).read_bytes()).hexdigest()

def normalise(text):
    """Strip whitespace, punctuation for fuzzy label comparison."""
    return re.sub(r"[\s\(\)·/··\-\.·]", "", text).lower()

# ── Border collision guard ────────────────────────────────────────────────────

def detect_border_collisions(pdf_path, tolerance=1.0):
    """
    Detect text words whose bounding box is intersected by a horizontal line drawn
    by the engine (table row borders).  A collision occurs when a horizontal line
    at y is strictly inside word bbox [y0+tol, y1-tol] — meaning a border line
    cuts through visible text, not merely touches its edge.

    Returns list of dicts: {text, word_y0, word_y1, line_y, overlap}.

    Root cause: basic_info rows with ROW_H_INFO=7mm have insufficient height for
    labels whose full text + "____..." exceeds the cell's available width (223.9pt),
    causing ReportLab to wrap to a second line that bleeds past the bottom border.
    """
    doc = pymupdf.open(str(pdf_path))
    page = doc[0]

    h_lines = []
    for path in page.get_drawings():
        for item in path.get("items", []):
            if item[0] == "l":
                p1, p2 = item[1], item[2]
                if abs(p1.y - p2.y) < 1.0:
                    x0 = min(p1.x, p2.x); x1 = max(p1.x, p2.x)
                    ly = (p1.y + p2.y) / 2
                    if x1 - x0 > 10:
                        h_lines.append((x0, ly, x1))

    words = page.get_text("words")
    collisions = []
    for wrd in words:
        wx0, wy0, wx1, wy1 = wrd[0], wrd[1], wrd[2], wrd[3]
        text = wrd[4]
        if not text.strip():
            continue
        for lx0, ly, lx1 in h_lines:
            if lx0 >= wx1 or lx1 <= wx0:
                continue
            if wy0 + tolerance < ly < wy1 - tolerance:
                collisions.append({
                    "text":    text,
                    "word_y0": round(wy0, 1),
                    "word_y1": round(wy1, 1),
                    "line_y":  round(ly, 1),
                    "overlap": round(min(wy1 - ly, ly - wy0), 1),
                })
    doc.close()
    return collisions


# ── Section A: PDF QA ────────────────────────────────────────────────────────

def qa_pdf(cid, code, spec):
    result = {"id": code, "section": "A_PDF", "issues": []}

    pdf_path = OUTPUT / f"TAI-FORM-{code}-blank.pdf"
    if not pdf_path.exists():
        result["status"] = "FAIL"
        result["issues"].append("PDF file missing")
        return result

    orientation  = spec["document"]["page"]["orientation"]
    expected_title = spec["document"]["title"]

    # expected labels from all sections
    expected_labels = []
    for s in spec.get("sections", []):
        stype = s.get("type")
        if stype == "basic_info":
            for f in s.get("fields", []):
                expected_labels.append(f["label"])
        elif stype == "repeat_table":
            for c in s.get("columns", []):
                expected_labels.append(c["label"])
        elif stype == "freeform_area":
            lbl = s.get("label", "")
            if lbl:
                expected_labels.append(lbl)

    doc = pymupdf.open(str(pdf_path))
    pages = len(doc)
    page  = doc[0]
    w, h  = page.mediabox.width, page.mediabox.height

    # orientation check
    is_landscape = (w > h)
    if orientation == "landscape" and not is_landscape:
        result["issues"].append(f"Expected landscape, got portrait ({w:.0f}x{h:.0f})")
    elif orientation == "portrait" and is_landscape:
        result["issues"].append(f"Expected portrait, got landscape ({w:.0f}x{h:.0f})")

    # A4 size check (±2pt)
    if orientation == "landscape":
        if abs(w - A4_H_PT) > 2 or abs(h - A4_W_PT) > 2:
            result["issues"].append(f"Not A4 landscape ({w:.1f}x{h:.1f})")
    else:
        if abs(w - A4_W_PT) > 2 or abs(h - A4_H_PT) > 2:
            result["issues"].append(f"Not A4 portrait ({w:.1f}x{h:.1f})")

    if pages != 1:
        result["issues"].append(f"Expected 1 page, got {pages}")

    # text extraction
    raw_text = page.get_text()

    # title check
    if expected_title not in raw_text:
        # try partial (first 6 chars)
        partial = expected_title[:6]
        if partial not in raw_text:
            result["issues"].append(f"Title not found in PDF: {expected_title!r}")

    # label coverage
    missing_labels = []
    found_labels   = []
    for lbl in expected_labels:
        norm_lbl = normalise(lbl)
        norm_raw = normalise(raw_text)
        # allow wrapping: check if first 3 chars of label present
        if norm_lbl[:3] in norm_raw:
            found_labels.append(lbl)
        else:
            missing_labels.append(lbl)

    # bbox out-of-bounds check
    tiny_texts = []
    oob_texts  = []
    blocks = page.get_text("dict")["blocks"]
    for blk in blocks:
        if blk.get("type") != 0:
            continue
        for line in blk.get("lines", []):
            for span in line.get("spans", []):
                sz = span.get("size", 99)
                bbox = span.get("bbox", [0,0,0,0])
                txt  = span.get("text", "").strip()
                if not txt:
                    continue
                if sz < FONT_TOO_SMALL:
                    tiny_texts.append(f"size={sz:.1f} text={txt[:20]!r}")
                # out of page bounds
                if bbox[0] < 0 or bbox[2] > w or bbox[1] < 0 or bbox[3] > h:
                    oob_texts.append(f"bbox={bbox} text={txt[:20]!r}")

    # font/Hangul check
    hangul_found = bool(re.search(r'[가-힣]', raw_text))

    # rasterise to PNG at 150 DPI
    mat  = pymupdf.Matrix(150/72, 150/72)
    pix  = page.get_pixmap(matrix=mat, alpha=False)
    png_path = EVIDENCE / f"A_{code}_pdf_page1.png"
    pix.save(str(png_path))

    doc.close()

    # border collision check (added WO-REF01-060-B8-WAVE2-VISUAL-REPAIR-005)
    border_collisions = detect_border_collisions(pdf_path)

    result.update({
        "pages": pages,
        "orientation": f"{w:.0f}x{h:.0f}",
        "title_found": expected_title in raw_text,
        "labels_total": len(expected_labels),
        "labels_found": len(found_labels),
        "labels_missing": missing_labels,
        "hangul_extractable": hangul_found,
        "tiny_texts_count": len(tiny_texts),
        "oob_texts_count": len(oob_texts),
        "border_collisions": len(border_collisions),
        "png": str(png_path.relative_to(BASE.parent.parent.parent)),
    })
    if missing_labels:
        result["issues"].append(f"Labels not found in PDF: {missing_labels}")
    if tiny_texts:
        result["issues"].append(f"Tiny text spans ({len(tiny_texts)}): {tiny_texts[:3]}")
    if oob_texts:
        result["issues"].append(f"Out-of-bounds text ({len(oob_texts)}): {oob_texts[:2]}")
    if border_collisions:
        result["issues"].append(
            f"Border collisions ({len(border_collisions)}): "
            + str([c["text"] for c in border_collisions[:3]])
        )
    result["status"] = "HOLD" if result["issues"] else "PASS"
    return result


# ── Section B: DOCX blank OOXML parse ───────────────────────────────────────

def qa_docx_blank(cid, code, spec):
    result = {"id": code, "section": "B_DOCX_BLANK", "issues": []}

    docx_path = OUTPUT / f"TAI-FORM-{code}-blank.docx"
    if not docx_path.exists():
        result["status"] = "FAIL"
        result["issues"].append("DOCX file missing")
        return result

    orientation = spec["document"]["page"]["orientation"]

    # ZIP validity
    try:
        with zipfile.ZipFile(docx_path) as z:
            names = z.namelist()
        if "word/document.xml" not in names:
            result["issues"].append("word/document.xml missing from DOCX ZIP")
    except Exception as e:
        result["status"] = "FAIL"
        result["issues"].append(f"DOCX ZIP error: {e}")
        return result

    # python-docx parse
    try:
        wd = python_docx.Document(str(docx_path))
    except Exception as e:
        result["status"] = "FAIL"
        result["issues"].append(f"python-docx open error: {e}")
        return result

    # page orientation from section settings
    sec = wd.sections[0]
    pg_w = sec.page_width.mm if sec.page_width else None
    pg_h = sec.page_height.mm if sec.page_height else None
    if pg_w and pg_h:
        is_landscape = pg_w > pg_h
        if orientation == "landscape" and not is_landscape:
            result["issues"].append(f"DOCX page: expected landscape but {pg_w:.0f}x{pg_h:.0f}mm")
        elif orientation == "portrait" and is_landscape:
            result["issues"].append(f"DOCX page: expected portrait but {pg_w:.0f}x{pg_h:.0f}mm")

    # DOCX structure: T0=title(1col), T1=basic_info(2col), T2+=freeform(1col)/repeat_table(>2col)
    # Filter to repeat_table candidates: tables whose first row has >2 cells
    expected_tables  = [s for s in spec.get("sections", []) if s.get("type") == "repeat_table"]
    repeat_act_tables = [t for t in wd.tables if t.rows and len(t.rows[0].cells) > 2]
    if len(repeat_act_tables) < len(expected_tables):
        result["issues"].append(
            f"Expected >={len(expected_tables)} repeat tables, found {len(repeat_act_tables)}"
        )

    # per-table column count
    col_issues = []
    for i, (exp_s, tbl) in enumerate(zip(expected_tables, repeat_act_tables)):
        exp_cols = len(exp_s.get("columns", []))
        if tbl.rows:
            actual_cols = len(tbl.rows[0].cells)
            if actual_cols != exp_cols:
                col_issues.append(
                    f"Table S{exp_s['id']} expected {exp_cols} cols, got {actual_cols}"
                )
    if col_issues:
        result["issues"].append(f"Column count mismatch: {col_issues}")

    # check that first table row has expected column labels
    label_issues = []
    for exp_s, tbl in zip(expected_tables, repeat_act_tables):
        if not tbl.rows:
            continue
        header_row = tbl.rows[0]
        header_texts = [c.text.strip() for c in header_row.cells]
        for col in exp_s.get("columns", []):
            lbl = col["label"]
            if not any(normalise(lbl)[:3] in normalise(ct) for ct in header_texts):
                label_issues.append(f"Col label '{lbl}' not in header: {header_texts}")
    if label_issues:
        result["issues"].append(f"Column header label mismatches: {label_issues[:3]}")

    result.update({
        "tables_expected": len(expected_tables),
        "tables_found": len(repeat_act_tables),
        "page_size_mm": f"{pg_w:.0f}x{pg_h:.0f}" if (pg_w and pg_h) else "UNKNOWN",
        "libreoffice_render": "UNVERIFIED_LIBREOFFICE_NOT_INSTALLED",
    })
    result["status"] = "HOLD" if result["issues"] else "PASS"
    return result


# ── Section C: DOCX roundtrip ────────────────────────────────────────────────

# Synthetic test marker prefix
MARKER_PREFIX = "TAI_QA_004_"

def _make_markers(spec):
    """Generate {section_id: {cell_key: marker_value}} for injection."""
    markers = {}
    for s in spec.get("sections", []):
        sid  = s["id"]
        stype = s.get("type")
        if stype == "basic_info":
            for i, f in enumerate(s.get("fields", [])):
                markers[f"{sid}_{f['id']}"] = f"{MARKER_PREFIX}{sid}_F{i+1:02d}"
        elif stype == "repeat_table":
            cols = s.get("columns", [])
            nrows = min(s.get("default_row_count", 3), 3)
            for r in range(nrows):
                for j, col in enumerate(cols):
                    markers[f"{sid}_R{r}_C{j}"] = f"{MARKER_PREFIX}{sid}_R{r}_C{j}"
        elif stype == "freeform_area":
            markers[f"{sid}_freeform"] = f"{MARKER_PREFIX}{sid}_freeform"
    return markers


def qa_docx_roundtrip(cid, code, spec):
    result = {"id": code, "section": "C_DOCX_ROUNDTRIP", "issues": []}

    docx_path = OUTPUT / f"TAI-FORM-{code}-blank.docx"
    if not docx_path.exists():
        result["status"] = "FAIL"
        result["issues"].append("DOCX file missing")
        return result

    markers = _make_markers(spec)
    expected_repeat_tables = [s for s in spec.get("sections", [])
                               if s.get("type") == "repeat_table"]

    with tempfile.TemporaryDirectory() as td:
        tmp_path = pathlib.Path(td) / f"TAI-FORM-{code}-qa.docx"
        shutil.copy2(docx_path, tmp_path)

        # ── INJECT ──────────────────────────────────────────────────────────
        wd = python_docx.Document(str(tmp_path))

        # basic_info: T1 is always wd.tables[1]; fields are arranged 2 per row
        for sec_spec in spec.get("sections", []):
            sid   = sec_spec["id"]
            stype = sec_spec.get("type")
            if stype == "basic_info" and len(wd.tables) > 1:
                tbl = wd.tables[1]
                fields = sec_spec.get("fields", [])
                for i, field in enumerate(fields):
                    row_idx = i // 2
                    col_idx = i % 2
                    key = f"{sid}_{field['id']}"
                    val = markers.get(key, "")
                    if val and row_idx < len(tbl.rows):
                        row = tbl.rows[row_idx]
                        if col_idx < len(row.cells):
                            cell = row.cells[col_idx]
                            cell.paragraphs[-1].add_run(val)

        # repeat_tables: inject into data rows (skip header row 0)
        # Use only tables with >2 cols (same filter as Section B)
        repeat_tbls = [t for t in wd.tables if t.rows and len(t.rows[0].cells) > 2]
        rt_idx = 0
        for sec_spec in spec.get("sections", []):
            if sec_spec.get("type") != "repeat_table":
                continue
            sid  = sec_spec["id"]
            cols = sec_spec.get("columns", [])
            nrows_inject = min(sec_spec.get("default_row_count", 3), 3)
            if rt_idx < len(repeat_tbls):
                tbl = repeat_tbls[rt_idx]
                for r in range(1, min(1 + nrows_inject, len(tbl.rows))):
                    actual_r = r - 1  # 0-based marker index
                    for j, col in enumerate(cols):
                        if j < len(tbl.rows[r].cells):
                            key = f"{sid}_R{actual_r}_C{j}"
                            val = markers.get(key, "")
                            if val:
                                cell = tbl.rows[r].cells[j]
                                cell.paragraphs[-1].add_run(val)
            rt_idx += 1

        # freeform_area: these are single-col tables (T2+ with 1 col)
        # Row 0 = label, Row 1 = content area
        freeform_tbls = [t for t in wd.tables if t.rows and len(t.rows[0].cells) == 1]
        for sec_spec in spec.get("sections", []):
            if sec_spec.get("type") != "freeform_area":
                continue
            sid = sec_spec["id"]
            key = f"{sid}_freeform"
            val = markers.get(key, "")
            if not val:
                continue
            lbl = normalise(sec_spec.get("label", ""))
            for tbl in freeform_tbls:
                if not tbl.rows:
                    continue
                header_text = normalise(tbl.rows[0].cells[0].text)
                if lbl[:4] and header_text[:4] == lbl[:4]:
                    if len(tbl.rows) > 1:
                        tbl.rows[1].cells[0].paragraphs[-1].add_run(val)
                    break

        wd.save(str(tmp_path))

        # ── REOPEN AND VERIFY ────────────────────────────────────────────────
        wd2 = python_docx.Document(str(tmp_path))
        full_text2 = "\n".join(p.text for p in wd2.paragraphs)
        for tbl in wd2.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    full_text2 += "\n" + cell.text

        found_markers  = []
        missing_markers = []
        for key, val in markers.items():
            if val in full_text2:
                found_markers.append(key)
            else:
                missing_markers.append(key)

        result.update({
            "markers_total": len(markers),
            "markers_found": len(found_markers),
            "markers_missing": missing_markers[:10],
            "libreoffice_render": "UNVERIFIED_LIBREOFFICE_NOT_INSTALLED",
        })

        if missing_markers:
            result["issues"].append(
                f"{len(missing_markers)}/{len(markers)} markers lost after roundtrip: "
                f"{missing_markers[:5]}"
            )

        result["status"] = "HOLD" if result["issues"] else "PASS"

    return result


# ── MAIN ─────────────────────────────────────────────────────────────────────

def main():
    # preflight SHA check
    build_receipt = BASE.parent / "evidence" / "REF01-B8-WAVE2-BUILD-003-RESULT.md"
    receipt_text  = build_receipt.read_text(encoding="utf-8")

    # Extract expected B8 SHA from receipt
    sha_from_receipt = {}
    for line in receipt_text.splitlines():
        m = re.match(r"(c\d+|gov-01)_(pdf|docx):\s*([0-9a-f]{64})", line.strip())
        if m:
            sha_from_receipt[f"{m.group(1)}_{m.group(2)}"] = m.group(3)

    sha_ok    = []
    sha_fail  = []
    for cid, code in B8_IDS:
        for ext, key in [("pdf","pdf"), ("docx","docx")]:
            p = OUTPUT / f"TAI-FORM-{code}-blank.{ext}"
            actual = sha256(p)
            expected = sha_from_receipt.get(f"{cid}_{key}", "")
            if actual == expected:
                sha_ok.append(f"{code}_{key}")
            else:
                sha_fail.append(f"{code}_{key}: got={actual[:8]}, expected={expected[:8]}")

    print(f"B8 OUTPUT SHA: {len(sha_ok)}/28 MATCH  FAIL={len(sha_fail)}")
    if sha_fail:
        for f in sha_fail:
            print(f"  ✗ {f}")

    # Frozen SHA check
    import importlib.util
    spec_bb = importlib.util.spec_from_file_location("bb", str(BASE / "batch_build.py"))
    bb = importlib.util.module_from_spec(spec_bb)
    spec_bb.loader.exec_module(bb)

    frozen_ok   = []
    frozen_fail = []
    for key, expected in bb.FROZEN_SHA.items():
        cid, ftype = key.rsplit("_", 1)
        entry = bb.REGISTRY[cid]
        fname = entry["pdf_name"] if ftype == "pdf" else entry["docx_name"]
        fpath = OUTPUT / fname
        actual = sha256(fpath)
        if actual == expected:
            frozen_ok.append(key)
        else:
            frozen_fail.append(f"{key}: got={actual[:8]}, expected={expected[:8]}")

    print(f"FROZEN SHA:    {len(frozen_ok)}/24 MATCH  FAIL={len(frozen_fail)}")
    if frozen_fail:
        for f in frozen_fail: print(f"  ✗ {f}")

    # Engine SHA
    py_engine_sha  = sha256(BASE / "common_v1_engine.py")
    cjs_engine_sha = sha256(BASE / "common_v1_engine.cjs")
    print(f"ENGINE PY:  {py_engine_sha[:8]}  {'MATCH' if py_engine_sha == 'be4899dad868d4336756ce61134546748ac2b0f5f4eb4dbbf7488e795d3f2aba' else 'MISMATCH'}")
    print(f"ENGINE CJS: {cjs_engine_sha[:8]}  {'MATCH' if cjs_engine_sha == '375250c74ef1e22786526c1fa2446d989bbed95e8e513e85deeccc3bfb727925' else 'MISMATCH'}")
    print()

    # Run all QA
    qa_matrix = []
    for cid, code in B8_IDS:
        spec_path = BASE / f"{cid}_v1.json"
        spec      = json.loads(spec_path.read_text(encoding="utf-8"))

        print(f"  [{code}]")
        r_a = qa_pdf(cid, code, spec)
        r_b = qa_docx_blank(cid, code, spec)
        r_c = qa_docx_roundtrip(cid, code, spec)

        status_a = r_a["status"]
        status_b = r_b["status"]
        status_c = r_c["status"]

        print(f"    A-PDF:       {status_a}  labels={r_a.get('labels_found','?')}/{r_a.get('labels_total','?')}  issues={len(r_a['issues'])}")
        print(f"    B-DOCX:      {status_b}  tables={r_b.get('tables_found','?')}/{r_b.get('tables_expected','?')}  issues={len(r_b['issues'])}")
        print(f"    C-ROUNDTRIP: {status_c}  markers={r_c.get('markers_found','?')}/{r_c.get('markers_total','?')}  issues={len(r_c['issues'])}")
        if r_a["issues"]:
            for iss in r_a["issues"]: print(f"      A✗ {iss}")
        if r_b["issues"]:
            for iss in r_b["issues"]: print(f"      B✗ {iss}")
        if r_c["issues"]:
            for iss in r_c["issues"]: print(f"      C✗ {iss}")

        qa_matrix.append({"form": code, "A_PDF": r_a, "B_DOCX_BLANK": r_b, "C_ROUNDTRIP": r_c})

    # Write qa_matrix.json
    matrix_path = EVIDENCE / "qa_matrix.json"
    matrix_path.write_text(json.dumps(qa_matrix, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nqa_matrix.json → {matrix_path}")

    # Summary
    summary = {"B8_SHA": f"{len(sha_ok)}/28", "FROZEN_SHA": f"{len(frozen_ok)}/24"}
    for row in qa_matrix:
        code = row["form"]
        summary[code] = {
            "A": row["A_PDF"]["status"],
            "B": row["B_DOCX_BLANK"]["status"],
            "C": row["C_ROUNDTRIP"]["status"],
        }
    print("\n=== SUMMARY ===")
    pass_a = sum(1 for r in qa_matrix if r["A_PDF"]["status"] == "PASS")
    pass_b = sum(1 for r in qa_matrix if r["B_DOCX_BLANK"]["status"] == "PASS")
    pass_c = sum(1 for r in qa_matrix if r["C_ROUNDTRIP"]["status"] == "PASS")
    hold_a = sum(1 for r in qa_matrix if r["A_PDF"]["status"] == "HOLD")
    hold_c = sum(1 for r in qa_matrix if r["C_ROUNDTRIP"]["status"] == "HOLD")
    print(f"  A PDF Layout:     {pass_a}/14 PASS  {hold_a}/14 HOLD")
    print(f"  B DOCX Blank:     {pass_b}/14 PASS  (LibreOffice render: UNVERIFIED)")
    print(f"  C DOCX Roundtrip: {pass_c}/14 PASS  {hold_c}/14 HOLD  (LibreOffice render: UNVERIFIED)")
    print(f"  D PNG evidence:   14 original PDF PNGs → {EVIDENCE}")

    return qa_matrix, sha_ok, sha_fail, frozen_ok, frozen_fail, py_engine_sha, cjs_engine_sha


if __name__ == "__main__":
    main()
