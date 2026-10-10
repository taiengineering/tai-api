"""
WO-REF01-060-B8-WAVE2-VISUAL-REPAIR-005-PHASE2-CANONICAL-006
One-shot protected canonical replacement for C031 / C043 / C044.
Run from any directory; uses absolute paths derived from __file__.
"""
import copy, hashlib, json, pathlib, shutil, subprocess, sys, tempfile
import pymupdf
import docx as python_docx

BASE     = pathlib.Path(__file__).parent
OUTPUT   = BASE.parent / "output"
EVIDENCE = BASE.parent / "evidence" / "visual-repair-005"
REPO_ROOT = BASE.parent.parent.parent  # tai-api root (.git lives here)

sys.path.insert(0, str(BASE))
import common_v1_engine as eng
eng.register_fonts()

# ── Constants ─────────────────────────────────────────────────────────────

EXPECTED_BASE_HEAD = "448c6f37a1e11504244d13446f48d450a1cc6bb1"

TARGETS = [("c031", "C031"), ("c043", "C043"), ("c044", "C044")]

ALLOWLIST_PATHS = {
    "c031": {
        "json":  BASE   / "c031_v1.json",
        "pdf":   OUTPUT / "TAI-FORM-C031-blank.pdf",
        "docx":  OUTPUT / "TAI-FORM-C031-blank.docx",
    },
    "c043": {
        "json":  BASE   / "c043_v1.json",
        "pdf":   OUTPUT / "TAI-FORM-C043-blank.pdf",
        "docx":  OUTPUT / "TAI-FORM-C043-blank.docx",
    },
    "c044": {
        "json":  BASE   / "c044_v1.json",
        "pdf":   OUTPUT / "TAI-FORM-C044-blank.pdf",
        "docx":  OUTPUT / "TAI-FORM-C044-blank.docx",
    },
}

EXPECTED_ORIG_SHA = {
    "c031_pdf":  "b12cd430e64f0938e908aa7b523ff1059799795414c2fd8c4c64ef6d185f5593",
    "c031_docx": "2f0e8d4431e816f5fa138dce36ff86b5d2932edffd753668f004980193206af5",
    "c043_pdf":  "275540374285c881beb0124e0393ff8e1acd3219bae71f22f603cacfcac065e4",
    "c043_docx": "f19ff42ab7e915a0d06ad796ca4fc595448410bdc735c671973f41022ca8a375",
    "c044_pdf":  "e7dc0b22ab666ff512a820aa235bcbffc0c68c89927ef9ce79d0137be47e80ca",
    "c044_docx": "d08e8fae4e75e13b724d978d0627670e7b7cd60fb3cda2403afbc83e630c5f1b",
}

EXPECTED_ENGINE_SHA = {
    "common_v1_engine.py":  "be4899dad868d4336756ce61134546748ac2b0f5f4eb4dbbf7488e795d3f2aba",
    "common_v1_engine.cjs": "375250c74ef1e22786526c1fa2446d989bbed95e8e513e85deeccc3bfb727925",
}

# Candidate JSON required labels
EXPECTED_CANDIDATE_LABELS = {
    "c031": {"S01": {"F06": "교육자료·버전", "F07": "출석증빙·보관"}},
    "c043": {"S01": {"F05": "참여/명단참조"}},
    "c044": {"S01": {"F01": "물질명/제품명", "F04": "대상·적용 작업"}},
}

MARKER_PREFIX = "TAI_QA_P2_"

# design_gate_notes suffix to replace
OLD_GATE_SUFFIX = "GPT 검토 전 빌드 차단."
NEW_GATE_SUFFIX = "내부 POC 후보 검토 통과, 법률·권리·외부공개 검토는 별도 대기."


# ── Utilities ──────────────────────────────────────────────────────────────

def sha256(p: pathlib.Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def abort(msg: str):
    print(f"\nPHASE2_HOLD: {msg}")
    sys.exit(1)


def detect_border_collisions(pdf_path, tolerance=1.0):
    doc  = pymupdf.open(str(pdf_path))
    page = doc[0]
    h_lines = []
    for path in page.get_drawings():
        for item in path.get("items", []):
            if item[0] == "l":
                p1, p2 = item[1], item[2]
                if abs(p1.y - p2.y) < 1.0:
                    x0, x1 = min(p1.x, p2.x), max(p1.x, p2.x)
                    ly = (p1.y + p2.y) / 2
                    if x1 - x0 > 10:
                        h_lines.append((x0, ly, x1))
    words = page.get_text("words")
    collisions = []
    for w in words:
        wx0, wy0, wx1, wy1, text = w[0], w[1], w[2], w[3], w[4]
        if not text.strip():
            continue
        for lx0, ly, lx1 in h_lines:
            if lx0 >= wx1 or lx1 <= wx0:
                continue
            if wy0 + tolerance < ly < wy1 - tolerance:
                collisions.append({"text": text, "line_y": round(ly, 1)})
    doc.close()
    return collisions


def rasterize_png(pdf_path, out_png, dpi=150):
    doc  = pymupdf.open(str(pdf_path))
    pix  = doc[0].get_pixmap(matrix=pymupdf.Matrix(dpi/72, dpi/72))
    pix.save(str(out_png))
    doc.close()


def docx_roundtrip(docx_path, spec):
    with tempfile.TemporaryDirectory() as td:
        tmp = pathlib.Path(td) / "qa.docx"
        shutil.copy2(docx_path, tmp)
        wd = python_docx.Document(str(tmp))

        markers = {}

        # basic_info: T1, fields 2-per-row
        for sec in spec["sections"]:
            if sec.get("type") == "basic_info" and len(wd.tables) > 1:
                tbl = wd.tables[1]
                for i, field in enumerate(sec.get("fields", [])):
                    rid, cid2 = i // 2, i % 2
                    key = f"{sec['id']}_{field['id']}"
                    val = MARKER_PREFIX + key
                    markers[key] = val
                    if rid < len(tbl.rows) and cid2 < len(tbl.rows[rid].cells):
                        tbl.rows[rid].cells[cid2].paragraphs[-1].add_run(val)

        # repeat_table: >2-col tables
        rep_tbls = [t for t in wd.tables if t.rows and len(t.rows[0].cells) > 2]
        rti = 0
        for sec in spec["sections"]:
            if sec.get("type") != "repeat_table":
                continue
            if rti < len(rep_tbls):
                tbl  = rep_tbls[rti]
                cols = sec.get("columns", [])
                for r in range(1, min(4, len(tbl.rows))):
                    for j, col in enumerate(cols):
                        key = f"{sec['id']}_R{r-1}_C{j}"
                        val = MARKER_PREFIX + key
                        markers[key] = val
                        if j < len(tbl.rows[r].cells):
                            tbl.rows[r].cells[j].paragraphs[-1].add_run(val)
            rti += 1

        # freeform_area: 1-col tables, skip T0
        ff_tbls = [t for t in wd.tables if t.rows and len(t.rows[0].cells) == 1][1:]
        ffi = 0
        for sec in spec["sections"]:
            if sec.get("type") != "freeform_area":
                continue
            if ffi < len(ff_tbls):
                tbl = ff_tbls[ffi]
                key = f"{sec['id']}_content"
                val = MARKER_PREFIX + key
                markers[key] = val
                if len(tbl.rows) > 1:
                    tbl.rows[1].cells[0].paragraphs[-1].add_run(val)
            ffi += 1

        wd.save(str(tmp))
        wd2    = python_docx.Document(str(tmp))
        text2  = "\n".join(p.text for p in wd2.paragraphs)
        for t in wd2.tables:
            for row in t.rows:
                for cell in row.cells:
                    text2 += "\n" + cell.text

        found   = [k for k, v in markers.items() if v in text2]
        missing = [k for k, v in markers.items() if v not in text2]
        return len(markers), len(found), missing


# ── Section 1: PRE-FLIGHT ─────────────────────────────────────────────────

def preflight():
    print("=== PHASE 2 PRE-FLIGHT ===\n")

    # HEAD
    result = subprocess.run(
        ["git", "rev-parse", "HEAD"], capture_output=True, text=True, cwd=str(REPO_ROOT)
    )
    current_head = result.stdout.strip()
    if current_head != EXPECTED_BASE_HEAD:
        abort(f"STALE_HEAD_BLOCKED: expected {EXPECTED_BASE_HEAD}, got {current_head}")
    print(f"  HEAD: {current_head}  MATCH")

    # Clean working tree: block only modified/staged/deleted, not untracked new evidence files
    result = subprocess.run(
        ["git", "status", "--porcelain"], capture_output=True, text=True, cwd=str(REPO_ROOT)
    )
    dirty = [ln for ln in result.stdout.splitlines() if ln and not ln.startswith("??")]
    if dirty:
        abort(f"DIRTY_TREE: {chr(10).join(dirty)}")
    print("  Working tree: CLEAN")

    # Engine SHAs
    for fname, exp in EXPECTED_ENGINE_SHA.items():
        actual = sha256(BASE / fname)
        if actual != exp:
            abort(f"ENGINE_SHA_MISMATCH: {fname}")
        print(f"  Engine {fname}: MATCH")

    # Original output SHAs
    for cid, _ in TARGETS:
        for ext in ("pdf", "docx"):
            key  = f"{cid}_{ext}"
            path = ALLOWLIST_PATHS[cid][ext]
            actual = sha256(path)
            exp    = EXPECTED_ORIG_SHA[key]
            if actual != exp:
                abort(f"ORIG_SHA_MISMATCH: {path.name} expected {exp[:16]}... got {actual[:16]}...")
            print(f"  Orig {path.name}: MATCH")

    # Candidate JSONs
    for cid, _ in TARGETS:
        cand_path = EVIDENCE / f"{cid}_candidate_v1.json"
        if not cand_path.exists():
            abort(f"CANDIDATE_MISSING: {cand_path}")
        spec = json.loads(cand_path.read_text(encoding="utf-8"))
        # Check labels
        for sid, fields in EXPECTED_CANDIDATE_LABELS.get(cid, {}).items():
            sec = next((s for s in spec["sections"] if s["id"] == sid), None)
            if sec is None:
                abort(f"CANDIDATE_SECTION_MISSING: {cid} {sid}")
            for fid, exp_label in fields.items():
                fld = next((f for f in sec.get("fields", []) if f["id"] == fid), None)
                if fld is None or fld["label"] != exp_label:
                    abort(f"CANDIDATE_LABEL_WRONG: {cid} {sid}.{fid} expected {exp_label!r}")
        # Check C031 text_flow
        if cid == "c031":
            s04 = next((s for s in spec["sections"] if s["id"] == "S04"), None)
            if s04 is None or s04.get("type") != "text_flow":
                abort("CANDIDATE_C031_S04_MISSING")
        # Check candidate metadata
        if spec.get("_meta", {}).get("candidate_status") != "REVIEW_ONLY_NOT_CANONICAL":
            abort(f"CANDIDATE_META_WRONG: {cid}")
        print(f"  Candidate {cid}: LABELS_OK / META_OK")

    print("\nPRE-FLIGHT: PASS\n")


# ── Section 2: PREPARE NEW CANONICAL JSONs ────────────────────────────────

def prepare_canonical_json(cid: str) -> dict:
    """Load candidate, remove review-only meta, update design_gate_notes."""
    cand_path = EVIDENCE / f"{cid}_candidate_v1.json"
    spec = json.loads(cand_path.read_text(encoding="utf-8"))
    meta = spec["_meta"]

    # Remove candidate review-only attrs
    meta.pop("candidate_for_review", None)
    meta.pop("candidate_status", None)

    # Update design_gate_notes: replace old build-block suffix
    notes = meta.get("design_gate_notes", "")
    if OLD_GATE_SUFFIX in notes:
        meta["design_gate_notes"] = notes.replace(OLD_GATE_SUFFIX, NEW_GATE_SUFFIX)
    elif not NEW_GATE_SUFFIX in notes:
        meta["design_gate_notes"] = notes.rstrip() + " " + NEW_GATE_SUFFIX

    return spec


# ── Section 3–4: GENERATE + VALIDATE in tmp ───────────────────────────────

EXPECTED_MARKERS = {"c031": 23, "c043": 22, "c044": 21}

def generate_and_validate(tmpdir: pathlib.Path):
    """Returns dict of {cid: {pdf, docx, sha_pdf, sha_docx, spec}} on success."""
    results = {}
    print("=== TEMP GENERATION + VALIDATION ===\n")

    for cid, code in TARGETS:
        print(f"[{code}]")
        spec = prepare_canonical_json(cid)

        tmp_json = tmpdir / f"{cid}_canonical_new.json"
        tmp_pdf  = tmpdir / f"TAI-FORM-{code}-blank.pdf"
        tmp_docx = tmpdir / f"TAI-FORM-{code}-blank.docx"

        # Write JSON to tmp
        tmp_json.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")

        # Generate PDF
        try:
            eng.generate_from_dict(spec, str(tmp_pdf))
        except Exception as e:
            abort(f"PDF_GEN_FAIL {code}: {e}")

        # Generate DOCX via Node engine
        result = subprocess.run(
            ["node", "gen_c002_docx_common.cjs", str(tmp_json), str(tmp_docx)],
            capture_output=True, text=True, cwd=str(BASE),
        )
        if result.returncode != 0:
            abort(f"DOCX_GEN_FAIL {code}: {result.stderr.strip()}")

        # PDF validation
        doc = pymupdf.open(str(tmp_pdf))
        n_pages = doc.page_count
        pdf_text = doc[0].get_text("text")
        doc.close()

        if n_pages != 1:
            abort(f"PDF_PAGE_COUNT {code}: expected 1, got {n_pages}")
        print(f"  Pages: {n_pages}  OK")

        collisions = detect_border_collisions(tmp_pdf)
        if collisions:
            abort(f"BORDER_COLLISION {code}: {len(collisions)} collisions found")
        print(f"  Border collisions: 0  OK")

        # C031 guidance check
        if cid == "c031":
            if "작성 안내" not in pdf_text:
                abort("C031_GUIDANCE_MISSING_PDF")
            print("  C031 guidance in PDF: PRESENT")

        # DOCX roundtrip
        exp_markers = EXPECTED_MARKERS[cid]
        total, found, missing = docx_roundtrip(tmp_docx, spec)
        if total != exp_markers:
            abort(f"MARKER_COUNT {code}: expected {exp_markers}, got {total}")
        if missing:
            abort(f"MARKER_LOSS {code}: {len(missing)} markers lost: {missing[:5]}")
        print(f"  DOCX roundtrip: {found}/{total}  OK")

        # C031 guidance in DOCX
        if cid == "c031":
            wd = python_docx.Document(str(tmp_docx))
            docx_text = "\n".join(p.text for p in wd.paragraphs)
            if "작성 안내" not in docx_text:
                abort("C031_GUIDANCE_MISSING_DOCX")
            print("  C031 guidance in DOCX: PRESENT")

        results[cid] = {
            "code":     code,
            "spec":     spec,
            "tmp_json": tmp_json,
            "tmp_pdf":  tmp_pdf,
            "tmp_docx": tmp_docx,
            "sha_pdf":  sha256(tmp_pdf),
            "sha_docx": sha256(tmp_docx),
            "sha_json": sha256(tmp_json),
            "collisions": 0,
            "pages":     n_pages,
            "markers":   total,
        }
        print()

    print("VALIDATION: PASS\n")
    return results


# ── Section 5: SHA MANIFEST + PNG ─────────────────────────────────────────

def write_manifest_and_pngs(results: dict):
    manifest = {
        "base_head": EXPECTED_BASE_HEAD,
        "engine": {
            "common_v1_engine.py":  sha256(BASE / "common_v1_engine.py"),
            "common_v1_engine.cjs": sha256(BASE / "common_v1_engine.cjs"),
        },
        "replacements": {},
    }
    for cid, r in results.items():
        code = r["code"]
        orig_key_pdf  = f"{cid}_pdf"
        orig_key_docx = f"{cid}_docx"
        manifest["replacements"][cid] = {
            "json_orig":  sha256(ALLOWLIST_PATHS[cid]["json"]),
            "json_new":   r["sha_json"],
            "pdf_orig":   EXPECTED_ORIG_SHA[orig_key_pdf],
            "pdf_new":    r["sha_pdf"],
            "docx_orig":  EXPECTED_ORIG_SHA[orig_key_docx],
            "docx_new":   r["sha_docx"],
            "candidate_json": sha256(EVIDENCE / f"{cid}_candidate_v1.json"),
        }
        # Rasterize new PDF
        png_path = EVIDENCE / f"NEW_{code}_pdf_page1.png"
        rasterize_png(r["tmp_pdf"], png_path)
        print(f"  PNG rasterized: {png_path.name}")

    manifest_path = EVIDENCE / "phase2_sha_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"  Manifest written: {manifest_path.name}\n")
    return manifest


# ── Section 6: ATOMIC REPLACEMENT TRANSACTION ─────────────────────────────

def atomic_replace(results: dict, backup_dir: pathlib.Path):
    print("=== ATOMIC REPLACEMENT ===\n")

    # Backup all 9 originals
    backed_up = {}
    for cid, _ in TARGETS:
        for key in ("json", "pdf", "docx"):
            src  = ALLOWLIST_PATHS[cid][key]
            dst  = backup_dir / f"{cid}_{key}_ORIG{src.suffix}"
            shutil.copy2(src, dst)
            backed_up[f"{cid}_{key}"] = (src, dst)
    print("  Backups: 9/9 OK")

    # Replacement order: JSON → PDF → DOCX
    replaced = []
    try:
        for cid, r in results.items():
            for key, tmp_path in [("json", r["tmp_json"]), ("pdf", r["tmp_pdf"]), ("docx", r["tmp_docx"])]:
                canonical = ALLOWLIST_PATHS[cid][key]
                shutil.copy2(tmp_path, canonical)
                replaced.append(f"{cid}_{key}")
                print(f"  Replaced: {canonical.name}")
    except Exception as e:
        print(f"\nREPLACEMENT_FAILED: {e} — rolling back...")
        for fkey, (src, backup) in backed_up.items():
            try:
                shutil.copy2(backup, src)
                print(f"  Rollback: {src.name}")
            except Exception as rb_err:
                print(f"  ROLLBACK_FAIL {src.name}: {rb_err}")
                abort(f"RECOVERY_REQUIRED: rollback failed for {src.name}")
        abort(f"REPLACEMENT_FAILED_ROLLED_BACK: {e}")

    print(f"\n  {len(replaced)}/9 files replaced successfully")

    # Post-replacement SHA verification
    errors = []
    for cid, r in results.items():
        for key, exp_sha in [("pdf", r["sha_pdf"]), ("docx", r["sha_docx"])]:
            actual = sha256(ALLOWLIST_PATHS[cid][key])
            if actual != exp_sha:
                errors.append(f"{cid}_{key}: sha mismatch after replacement")
    if errors:
        abort(f"POST_REPLACE_SHA_FAIL: {errors}")
    print("  Post-replacement SHA: 6/6 NEW_SHA_VERIFIED\n")


# ── Section 7: POST-REPLACEMENT VERIFICATION ──────────────────────────────

def post_verify(results: dict):
    print("=== POST-REPLACEMENT VERIFICATION ===\n")

    # Border collisions on actual canonical PDFs
    for cid, r in results.items():
        code = r["code"]
        canon_pdf = ALLOWLIST_PATHS[cid]["pdf"]
        colls = detect_border_collisions(canon_pdf)
        if colls:
            abort(f"POST_COLLISION {code}: {len(colls)} collisions in canonical PDF")
        print(f"  {code} canonical PDF collisions: 0  OK")

    # Engine SHAs unchanged
    for fname, exp in EXPECTED_ENGINE_SHA.items():
        actual = sha256(BASE / fname)
        if actual != exp:
            abort(f"ENGINE_SHA_CHANGED: {fname}")
    print("  Engine SHAs: 2/2 MATCH\n")


# ── Section 8: phase2_qa_matrix.json ──────────────────────────────────────

def write_qa_matrix(results: dict):
    matrix = {}
    for cid, r in results.items():
        matrix[r["code"]] = {
            "pages":              r["pages"],
            "border_collisions":  r["collisions"],
            "roundtrip_markers":  r["markers"],
            "roundtrip_status":   "PASS",
            "guidance_pdf":       "PRESENT" if cid == "c031" else "N/A",
            "guidance_docx":      "PRESENT" if cid == "c031" else "N/A",
            "docx_libreoffice":   "UNVERIFIED_LIBREOFFICE_NOT_INSTALLED",
            "sha_pdf_new":        r["sha_pdf"],
            "sha_docx_new":       r["sha_docx"],
            "status":             "PASS",
        }
    out = EVIDENCE / "phase2_qa_matrix.json"
    out.write_text(json.dumps(matrix, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"  QA matrix: {out.name}\n")


# ── MAIN ───────────────────────────────────────────────────────────────────

def main():
    print("=" * 60)
    print("WO-REF01-060-B8-WAVE2-VISUAL-REPAIR-005-PHASE2-CANONICAL-006")
    print("=" * 60 + "\n")

    preflight()

    with tempfile.TemporaryDirectory() as td:
        tmpdir     = pathlib.Path(td)
        backup_dir = pathlib.Path(tempfile.mkdtemp(prefix="phase2_backup_"))

        results = generate_and_validate(tmpdir)

        print("=== SHA MANIFEST + PNGs ===\n")
        manifest = write_manifest_and_pngs(results)

        atomic_replace(results, backup_dir)
        post_verify(results)

        print("=== QA MATRIX ===\n")
        write_qa_matrix(results)

    # Final report
    print("=" * 60)
    print("PHASE 2 FINAL REPORT")
    print("=" * 60)
    for cid, r in results.items():
        code = r["code"]
        print(f"\n  {code}:")
        print(f"    pdf_new_sha8  = {r['sha_pdf'][:8]}")
        print(f"    docx_new_sha8 = {r['sha_docx'][:8]}")
        print(f"    collisions    = {r['collisions']}")
        print(f"    roundtrip     = {r['markers']}/{r['markers']}")

    print(f"""
WO                          = WO-REF01-060-B8-WAVE2-VISUAL-REPAIR-005-PHASE2-CANONICAL-006
BASE_HEAD                   = {EXPECTED_BASE_HEAD}
OWNER_PHASE2_AUTHORIZATION  = YES_LIMITED_3_FORMS
CANONICAL_JSON_CHANGED      = 3/3 EXACT_SCOPE
CANONICAL_PDF_REPLACED      = 3/3
CANONICAL_DOCX_REPLACED     = 3/3
BORDER_COLLISIONS           = C031:0 / C043:0 / C044:0
C031_GUIDANCE_PDF_DOCX      = PRESENT / PRESENT
DOCX_ROUNDTRIP              = 23/23 / 22/22 / 21/21
DOCX_VISUAL                 = UNVERIFIED_LIBREOFFICE_NOT_INSTALLED
COMMON_ENGINE_SHA           = 2/2 MATCH
DB_WRITES                   = 0
PR_MERGE                    = BLOCKED
PRODUCTION_DEPLOY           = BLOCKED
PUBLICATION                 = INTERNAL_POC_ONLY
LEGAL_REVIEW                = PENDING
RIGHTS                      = UNVERIFIED
GPT_INDEPENDENT_VERIFY      = PENDING
B8_CLOSED_FINAL             = NO
""")


if __name__ == "__main__":
    main()
