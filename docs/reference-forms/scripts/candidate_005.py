"""
WO-REF01-060-B8-WAVE2-VISUAL-REPAIR-005  Phase 1 / 005A / 005B
Candidate JSON generation, PDF/DOCX creation (temp only), border-collision detection,
150 DPI PNG rasterization, and DOCX roundtrip for C031/C043/C044.
"""
import copy, hashlib, json, pathlib, shutil, subprocess, sys, tempfile, re
import pymupdf
import docx as python_docx

BASE    = pathlib.Path(__file__).parent
OUTPUT  = BASE.parent / "output"
EVIDENCE = BASE.parent / "evidence" / "visual-repair-005"
EVIDENCE.mkdir(parents=True, exist_ok=True)

sys.path.insert(0, str(BASE))
import common_v1_engine as eng
eng.register_fonts()

MARKER_PREFIX = "TAI_QA_005_"

# ── candidate label corrections ────────────────────────────────────────────
CORRECTIONS = {
    "c031": {
        "S01": {
            "F06": "교육자료·버전",            # was "사용 교육자료(명칭·버전)"  217.73pt OK
            "F07": "출석증빙·보관",            # was "출석 증빙 종류·보관 위치"  217.73pt OK  (005A: 출석 보존)
        }
    },
    "c043": {
        "S01": {
            "F05": "참여/명단참조",            # was "훈련 참여자/명단 참조"     218.58pt OK  (005A: 참조 복원)
        }
    },
    "c044": {
        "S01": {
            "F01": "물질명/제품명",            # was "화학물질/제품명"           219.34pt OK
            "F04": "대상·적용 작업",           # was "작업대상 또는 적용 작업"  220.53pt OK  (005A: 두 개념 구분 유지)
        }
    },
}

TARGETS = [("c031","C031"), ("c043","C043"), ("c044","C044")]

# ── 005B: candidate metadata ───────────────────────────────────────────────
CANDIDATE_META = {
    "candidate_for_review": "WO-REF01-060-B8-WAVE2-VISUAL-REPAIR-005B",
    "candidate_status":     "REVIEW_ONLY_NOT_CANONICAL",
}

# ── 005B: C031 입력 안내 text_flow ─────────────────────────────────────────
C031_GUIDANCE_SECTION = {
    "id":   "S04",
    "type": "text_flow",
    "paragraphs": [
        {
            "id":    "P01",
            "text":  "【작성 안내】 출석증빙·보관란에는 출석 증빙의 종류(예: 참석자 서명부)와 실제 보관 위치를 함께 기재하십시오.",
            "align": "left",
        }
    ],
}


def sha256(p):
    return hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()


def apply_corrections(spec, cid):
    s = copy.deepcopy(spec)
    # Label corrections
    corr = CORRECTIONS.get(cid, {})
    for section in s["sections"]:
        sid = section.get("id")
        if sid in corr:
            field_corr = corr[sid]
            for field in section.get("fields", []):
                if field["id"] in field_corr:
                    field["label"] = field_corr[field["id"]]
    # 005B: add candidate metadata to _meta
    s["_meta"].update(CANDIDATE_META)
    # 005B: C031 only — append text_flow guidance at end of sections
    if cid == "c031":
        s["sections"].append(copy.deepcopy(C031_GUIDANCE_SECTION))
    return s


# ── Border collision detection ─────────────────────────────────────────────

def detect_border_collisions(pdf_path, tolerance=1.0):
    """
    Return list of dicts: words whose bbox is intersected by a horizontal line.
    Uses PyMuPDF page.get_drawings() for lines and page.get_text('words') for words.
    A 'collision' = horizontal border line y is strictly inside word bbox [y0+tol, y1-tol].
    """
    doc = pymupdf.open(str(pdf_path))
    page = doc[0]

    # Collect horizontal lines from page drawings
    h_lines = []
    for path in page.get_drawings():
        for item in path.get("items", []):
            if item[0] == "l":               # line segment
                p1, p2 = item[1], item[2]
                if abs(p1.y - p2.y) < 1.0:  # horizontal (within 1pt)
                    x0 = min(p1.x, p2.x)
                    x1 = max(p1.x, p2.x)
                    ly = (p1.y + p2.y) / 2
                    if x1 - x0 > 10:        # skip tiny decorative lines
                        h_lines.append((x0, ly, x1))

    # Check each word against horizontal lines
    words = page.get_text("words")  # (x0, y0, x1, y1, text, block, line, word)
    collisions = []
    for w in words:
        wx0, wy0, wx1, wy1 = w[0], w[1], w[2], w[3]
        text = w[4]
        if not text.strip():
            continue
        for lx0, ly, lx1 in h_lines:
            # x overlap
            if lx0 >= wx1 or lx1 <= wx0:
                continue
            # line cuts through middle of word bbox (not just edge)
            if wy0 + tolerance < ly < wy1 - tolerance:
                collisions.append({
                    "text":    text,
                    "word_y0": round(wy0, 1),
                    "word_y1": round(wy1, 1),
                    "line_y":  round(ly, 1),
                    "overlap":  round(min(wy1 - ly, ly - wy0), 1),
                })
    doc.close()
    return collisions


def rasterize_png(pdf_path, out_png, dpi=150):
    doc = pymupdf.open(str(pdf_path))
    page = doc[0]
    mat = pymupdf.Matrix(dpi / 72, dpi / 72)
    pix = page.get_pixmap(matrix=mat)
    pix.save(str(out_png))
    doc.close()


def docx_roundtrip(docx_path, cid, spec_cand):
    """
    Inject markers into basic_info, repeat_table, and freeform_area cells,
    save, reopen, verify. Returns (markers_total, markers_found, missing).
    """
    spec = spec_cand

    with tempfile.TemporaryDirectory() as td:
        tmp = pathlib.Path(td) / "qa.docx"
        shutil.copy2(docx_path, tmp)
        wd = python_docx.Document(str(tmp))

        markers = {}
        # basic_info: T1 (index 1), fields 2 per row
        for sec in spec["sections"]:
            if sec.get("type") == "basic_info" and len(wd.tables) > 1:
                tbl = wd.tables[1]
                for i, field in enumerate(sec.get("fields", [])):
                    rid = i // 2; cid2 = i % 2
                    key = f"{sec['id']}_{field['id']}"
                    val = f"{MARKER_PREFIX}{key}"
                    markers[key] = val
                    if rid < len(tbl.rows) and cid2 < len(tbl.rows[rid].cells):
                        tbl.rows[rid].cells[cid2].paragraphs[-1].add_run(val)

        # repeat_table: >2-col tables
        rep_tbls = [t for t in wd.tables if t.rows and len(t.rows[0].cells) > 2]
        rti = 0
        for sec in spec["sections"]:
            if sec.get("type") != "repeat_table":
                continue
            if rti < len(rep_tbls):
                tbl = rep_tbls[rti]
                cols = sec.get("columns", [])
                for r in range(1, min(4, len(tbl.rows))):
                    for j, col in enumerate(cols):
                        key = f"{sec['id']}_R{r-1}_C{j}"
                        val = f"{MARKER_PREFIX}{key}"
                        markers[key] = val
                        if j < len(tbl.rows[r].cells):
                            tbl.rows[r].cells[j].paragraphs[-1].add_run(val)
            rti += 1

        # freeform_area: 1-col tables (skip T0 title at index 0)
        ff_tbls = [t for t in wd.tables if t.rows and len(t.rows[0].cells) == 1][1:]
        ffi = 0
        for sec in spec["sections"]:
            if sec.get("type") != "freeform_area":
                continue
            if ffi < len(ff_tbls):
                tbl = ff_tbls[ffi]
                key = f"{sec['id']}_content"
                val = f"{MARKER_PREFIX}{key}"
                markers[key] = val
                if len(tbl.rows) > 1:
                    tbl.rows[1].cells[0].paragraphs[-1].add_run(val)
            ffi += 1

        wd.save(str(tmp))
        wd2 = python_docx.Document(str(tmp))
        text2 = "\n".join(p.text for p in wd2.paragraphs)
        for tbl in wd2.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    text2 += "\n" + cell.text

        found = [k for k, v in markers.items() if v in text2]
        missing = [k for k, v in markers.items() if v not in text2]
        return len(markers), len(found), missing


# ── Main ───────────────────────────────────────────────────────────────────

def main():
    print("=== WO-REF01-060-B8-WAVE2-VISUAL-REPAIR-005B Phase 1 Final ===\n")

    results = {}

    with tempfile.TemporaryDirectory() as td:
        tmpdir = pathlib.Path(td)

        for cid, code in TARGETS:
            print(f"[{code}]")

            # Load original spec
            spec_path = BASE / f"{cid}_v1.json"
            spec_orig = json.loads(spec_path.read_text(encoding="utf-8"))

            # Apply candidate corrections
            spec_cand = apply_corrections(spec_orig, cid)

            # Show before/after labels
            corr = CORRECTIONS.get(cid, {})
            for sid, fields in corr.items():
                for fid, new_lbl in fields.items():
                    orig_sec = next(s for s in spec_orig["sections"] if s["id"] == sid)
                    orig_fld = next(f for f in orig_sec.get("fields", []) if f["id"] == fid)
                    print(f"  {sid}.{fid}: '{orig_fld['label']}' → '{new_lbl}'")

            # Write candidate JSON to temp AND to evidence (REVIEW_ONLY_NOT_CANONICAL)
            cand_json = tmpdir / f"{cid}_candidate.json"
            payload = json.dumps(spec_cand, ensure_ascii=False, indent=2)
            cand_json.write_text(payload, encoding="utf-8")
            ev_json = EVIDENCE / f"{cid}_candidate_v1.json"
            ev_json.write_text(payload, encoding="utf-8")

            # Generate candidate PDF + DOCX using existing engine
            cand_pdf  = tmpdir / f"TAI-FORM-{code}-candidate.pdf"
            cand_docx = tmpdir / f"TAI-FORM-{code}-candidate.docx"

            try:
                # PDF via Python engine
                eng.generate_from_dict(spec_cand, str(cand_pdf))
                # DOCX via Node engine (write candidate JSON to temp, then invoke)
                result = subprocess.run(
                    ["node", "gen_c002_docx_common.cjs", str(cand_json), str(cand_docx)],
                    capture_output=True, text=True, cwd=str(BASE),
                )
                if result.returncode != 0:
                    raise RuntimeError(result.stderr.strip() or "node exited non-zero")
                print(f"  Engine: PDF+DOCX generated OK")
            except Exception as e:
                print(f"  ERROR generating outputs: {e}")
                results[code] = {"status": "ERROR", "error": str(e)}
                continue

            # Verify candidate is 1 page
            doc = pymupdf.open(str(cand_pdf))
            n_pages = doc.page_count
            doc.close()
            print(f"  Pages: {n_pages}")

            # Collision detection: ORIGINAL
            orig_pdf = OUTPUT / f"TAI-FORM-{code}-blank.pdf"
            orig_collisions = detect_border_collisions(orig_pdf)
            print(f"  Orig collisions: {len(orig_collisions)}")
            for c in orig_collisions[:3]:
                print(f"    '{c['text']}' y={c['word_y0']:.1f}-{c['word_y1']:.1f} border@{c['line_y']:.1f}")

            # Collision detection: CANDIDATE
            cand_collisions = detect_border_collisions(cand_pdf)
            print(f"  Cand collisions: {len(cand_collisions)}")

            # Rasterize original PDF (already exists in qa-004, but re-rasterize for comparison)
            orig_png = EVIDENCE / f"ORIG_{code}_pdf_page1.png"
            rasterize_png(orig_pdf, orig_png, dpi=150)

            # Rasterize candidate PDF
            cand_png = EVIDENCE / f"CAND_{code}_pdf_page1.png"
            rasterize_png(cand_pdf, cand_png, dpi=150)
            print(f"  PNGs: {orig_png.name}, {cand_png.name}")

            # DOCX roundtrip
            total, found, missing = docx_roundtrip(cand_docx, cid, spec_cand)
            roundtrip_status = "PASS" if not missing else f"FAIL({len(missing)}/{total} lost)"
            print(f"  DOCX roundtrip: {roundtrip_status}")

            # 005B: C031 — verify guidance text exists in PDF and DOCX
            guidance_pdf_ok  = None
            guidance_docx_ok = None
            if cid == "c031":
                guidance_snippet = "작성 안내"
                pdf_text = pymupdf.open(str(cand_pdf))[0].get_text("text")
                guidance_pdf_ok = guidance_snippet in pdf_text
                wd_check = python_docx.Document(str(cand_docx))
                all_text = "\n".join(p.text for p in wd_check.paragraphs)
                guidance_docx_ok = guidance_snippet in all_text
                print(f"  Guidance in PDF: {guidance_pdf_ok}  DOCX: {guidance_docx_ok}")

            # 005B: verify candidate metadata present
            meta_ok = (
                spec_cand.get("_meta", {}).get("candidate_status") == "REVIEW_ONLY_NOT_CANONICAL"
            )

            results[code] = {
                "orig_collisions": len(orig_collisions),
                "orig_collision_texts": [c["text"] for c in orig_collisions[:5]],
                "cand_collisions": len(cand_collisions),
                "pages": n_pages,
                "orig_png": str(orig_png.relative_to(BASE.parent.parent)),
                "cand_png": str(cand_png.relative_to(BASE.parent.parent)),
                "docx_roundtrip": roundtrip_status,
                "markers_total": total,
                "markers_found": found,
                "candidate_metadata": "PRESENT" if meta_ok else "MISSING",
                "guidance_pdf": ("PRESENT" if guidance_pdf_ok else "MISSING") if cid == "c031" else "N/A",
                "guidance_docx": ("PRESENT" if guidance_docx_ok else "MISSING") if cid == "c031" else "N/A",
                "label_corrections": {
                    f"{sid}.{fid}": {"from": next(f["label"] for s in spec_orig["sections"]
                                                   if s["id"]==sid
                                                   for f in s.get("fields",[]) if f["id"]==fid),
                                     "to": new_lbl}
                    for sid, fields in corr.items()
                    for fid, new_lbl in fields.items()
                },
                "status": "PASS" if (not cand_collisions and n_pages == 1 and not missing
                                     and meta_ok) else "HOLD",
            }

            print()

    # Write qa_matrix_005.json
    matrix_path = EVIDENCE / "qa_matrix_005.json"
    matrix_path.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"qa_matrix_005.json → {matrix_path}")

    print("\n=== SUMMARY ===")
    for code, r in results.items():
        status = r.get("status", "?")
        oc = r.get("orig_collisions", "?")
        cc = r.get("cand_collisions", "?")
        rt = r.get("docx_roundtrip", "?")
        print(f"  {code}: {status}  orig_collisions={oc} → cand_collisions={cc}  roundtrip={rt}")


if __name__ == "__main__":
    main()
